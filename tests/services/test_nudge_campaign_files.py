"""Campaigns in and out of Excel, CSV and JSON (services/nudges/campaign_files.py)."""

from __future__ import annotations

import io

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.db.session import get_session
from app.services.nudges import campaign_files as cf
from app.services.nudges import store


async def load(session, name: str, text: str, *, commit=True, mode=cf.UPSERT) -> dict:
    return await cf.import_rows(session, cf.read_file(name, text.encode()), mode=mode,
                                commit=commit, actor="ops@test")


async def by_name(session) -> dict:
    return {c.name: c for c in await store.list_campaigns(session)}


@pytest.mark.parametrize("fmt", ["xlsx", "csv", "json"])
async def test_the_example_file_imports_and_exports_back_unchanged(session, fmt):
    raw, _ = cf.sample(fmt)
    result = await cf.import_rows(session, cf.read_file(f"sample.{fmt}", raw), commit=True,
                                  actor="ops@test")
    assert result["counts"]["create"] == 3 and result["counts"]["error"] == 0, result["rows"]
    campaigns = await by_name(session)
    assert not any(c.is_active for c in campaigns.values())                 # all paused
    smo = campaigns["6-month follow-up (SMO)"].config["journey"]
    assert smo["no_response"]["steps"] == [{"every_days": 3, "times": 3},
                                           {"every_days": 10, "times": 3},
                                           {"every_days": 30, "times": None}]
    assert smo["no"]["max_attempts"] == 6 and smo["length_days"] == 183
    unpaid = campaigns["Unpaid order follow-up"]
    assert unpaid.kind == "abandoned_checkout" and unpaid.config["trigger"]["restore_cart"] is True
    assert unpaid.config["schedule"]["send_window"] == {"start": "10:00", "end": "19:00"}

    views = [store.campaign_view(c) for c in campaigns.values()]
    out, _ = cf.export(views, fmt)
    again = await cf.import_rows(session, cf.read_file(f"back.{fmt}", out), commit=False,
                                 actor="ops@test")
    assert again["counts"] == {"create": 0, "update": 0, "unchanged": 3, "skip": 0, "error": 0}, \
        again["rows"]


async def test_a_spreadsheet_row_sets_every_part_of_a_campaign(session):
    sheet = ("Name,Type,Days since first message,Message pool,Journey length (days),Days per stage,"
             "No response: follow-ups,Later: follow-ups,No: follow-ups,No: most sends,Send from,Send until,Send on,"
             "Don't repeat the angle of the last,Resume after an unfinished order (days)\n"
             'Kerala journey,Follow-up journey,2,KERALA,90,15,"2x4, 7x2, 21","5x2, 14",45,3,10:00,20:00,"mon, wed, fri",4,2.5\n')
    result = await load(session, "plan.csv", sheet)
    assert result["counts"]["create"] == 1, result["rows"]
    c = (await by_name(session))["Kerala journey"]
    j = c.config["journey"]
    assert (c.kind, c.config["trigger"]["days_since_signup"], j["pool"], j["length_days"],
            j["stage_days"]) == ("follow_up", 2, "KERALA", 90, 15)
    assert cf.format_steps(j["no_response"]["steps"]) == "2x4, 7x2, 21"
    assert cf.format_steps(j["later"]["steps"]) == "5x2, 14"
    assert (cf.format_steps(j["no"]["steps"]), j["no"]["max_attempts"]) == ("45", 3)
    assert c.config["schedule"]["days"] == ["mon", "wed", "fri"]
    assert (j["avoid_angle_repeat"], j["resume_after_days"]) == (4, 2.5)


async def test_updating_by_name_changes_only_the_cells_given(session):
    await load(session, "a.csv", "Name,Type,Wait (minutes),Give up after (hours),Days before again\n"
                                 "Win-back,Abandoned cart,30,24,7\n")
    c = (await by_name(session))["Win-back"]
    text = c.config["message"]["text"]

    check = await load(session, "b.csv", "Name,Days before again,Max per customer\n"
                                         "win-back,14,2\n", commit=False)
    assert check["counts"]["update"] == 1
    assert sorted(check["rows"][0]["changes"]) == ["Days before again", "Max per customer"]
    assert c.config["limits"]["cooldown_days"] == 7                          # only checked

    await load(session, "b.csv", "Name,Days before again,Max per customer\nwin-back,14,2\n")
    assert (c.config["limits"]["cooldown_days"], c.config["limits"]["max_per_customer"]) == (14, 2)
    assert c.config["trigger"]["after_minutes"] == 30 and c.config["message"]["text"] == text
    assert c.name == "Win-back"                                              # not renamed

    # "-" clears a setting; add-only leaves an existing campaign alone.
    await load(session, "c.csv", "Name,Max per customer\nWin-back,-\n")
    assert c.config["limits"]["max_per_customer"] is None
    skipped = await load(session, "d.csv", "Name,Type,Wait (minutes)\nWin-back,Abandoned cart,90\n",
                         mode=cf.ADD_ONLY)
    assert skipped["counts"]["skip"] == 1 and c.config["trigger"]["after_minutes"] == 30


async def test_rows_with_problems_are_skipped_with_the_reason(session):
    sheet = ("Name,Type,Wait (minutes),Give up after (hours),No response: follow-ups,Message\n"
             "Good cart,Abandoned cart,60,23,,Hi {first_name}\n"
             "Bad type,Something else,60,23,,Hi\n"
             "Bad wait,Abandoned cart,3000,23,,Hi\n"
             "Bad steps,Follow-up journey,,,every now and then,Hi\n"
             "No type,,,,,Hi\n"
             "Retired,Repeat purchase,,,,Hi\n"
             "Journey,Follow-up journey,,,,\n"
             "Good cart,Abandoned cart,30,23,,Again\n")
    result = await load(session, "x.csv", sheet)
    rows = {(r["row"], r["code"]): r for r in result["rows"]}
    assert result["counts"]["create"] == 2 and result["counts"]["error"] == 6
    reasons = {name: r["error"] for (_, name), r in rows.items() if r["action"] == "error"}
    assert "is not one of" in reasons["Bad type"]
    assert "after_minutes" in reasons["Bad wait"]
    assert "every now and then" in reasons["Bad steps"]
    assert "Type is needed" in reasons["No type"]
    # A type the page no longer offers is refused.
    assert "is not one of" in reasons["Retired"]
    assert "Same name as row 2" in rows[(9, "Good cart")]["error"]
    campaigns = await by_name(session)
    assert set(campaigns) == {"Good cart", "Journey"}
    assert campaigns["Good cart"].config["trigger"]["after_minutes"] == 60


async def test_files_that_cannot_be_read(session):
    for name, raw, words in (("x.json", b"{oops", "not valid JSON"),
                             ("x.csv", b"Title,Message\nA,B\n", "Name column"),
                             ("x.xls", b"\xd0\xcf\x11\xe0 old excel", ".xls"),
                             ("x.csv", b"", "empty")):
        with pytest.raises(cf.NudgeConfigError, match=words):
            cf.read_file(name, raw)


@pytest.fixture
def api(session):
    from app.api.routes.nudges import nudge_actor, router
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_session] = lambda: session
    app.dependency_overrides[nudge_actor] = lambda: "ops@test"
    return TestClient(app)


async def test_the_api_and_the_page(api, session):
    sample = api.get("/api/nudges/campaigns/sample", params={"format": "xlsx"})
    assert sample.status_code == 200 and sample.content[:2] == b"PK"
    check = api.post("/api/nudges/campaigns/import",
                     files={"file": ("sample.xlsx", sample.content)}).json()
    assert check["committed"] is False and check["counts"]["create"] == 3
    assert api.get("/api/nudges/campaigns").json()["campaigns"] == []
    done = api.post("/api/nudges/campaigns/import", files={"file": ("sample.xlsx", sample.content)},
                    data={"commit": "true"}).json()
    assert done["counts"]["create"] == 3
    listing = api.get("/api/nudges/campaigns").json()
    assert len(listing["campaigns"]) == 3 and listing["columns"][0]["heading"] == "Name"

    exported = api.get("/api/nudges/campaigns/export", params={"format": "json"}).json()
    assert {c["type"] for c in exported["campaigns"]} == {"Follow-up journey", "Abandoned cart",
                                                          "Unpaid order"}
    assert all("audience" not in c["config"] and "offer" not in c["config"]
               for c in exported["campaigns"])
    assert api.get("/api/nudges/campaigns/export", params={"format": "pdf"}).status_code == 422
    assert "campaign_columns" in api.get("/api/nudges/reference").json()

    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(api.get("/api/nudges/campaigns/export").content))
    assert book.sheetnames == ["Campaigns", "How to fill it in"]
    headings = [c.value for c in book["Campaigns"][1]]
    assert headings[:2] == ["Name", "Type"]
    assert not {"Priority", "Coupon code"} & set(headings)
    assert "No response: follow-ups" in headings

    from pathlib import Path
    page = Path("app/templates/admin/nudges.html").read_text(encoding="utf-8")
    assert "STARTERS" not in page and "data-starter" not in page            # no fixed cards
    assert "openCampaignImport" in page and "/campaigns/export" in page
