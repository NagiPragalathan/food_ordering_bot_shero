"""The message library a follow-up journey picks from - and its import/export.

Messages are rows the team adds on the Nudges page (Messages tab), or brings
in all at once from Excel, CSV or JSON. An import is matched on the message
ID ("M1-01"): a new ID is added, a known one is updated. Every file is
checked first and shown row by row, so nothing changes until the team says
so, and a row with a problem is skipped with the reason.

The columns are the same in every format, and an export is a file that
imports straight back:

    ID, Pool, Month, Title, Angle, Message, Groups, Occasion, From, Until,
    Order by, Kitchens, Dishes, WhatsApp template, Template values,
    Template button, Active

Only ID, Title and Message are needed for a new message. On an update only
the columns in the file change, so a sheet with just ID and Active switches
messages on and off.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timezone

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.logging import get_logger
from app.db.models import JourneyBucket, NudgeMessage
from app.services.nudges.config import (
    BUCKET_LABELS,
    MESSAGE_MAX,
    NudgeConfigError,
    check_library_text,
    check_placeholders,
)

log = get_logger(__name__)

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_ROWS = 2000
CODE = re.compile(r"^[A-Z0-9][A-Z0-9 _.\-/]{0,39}$")
TEMPLATE_NAME = re.compile(r"^[a-z0-9_]+$")
BUTTONS = ("none", "order_link")
DEFAULT_POOL = "SMO"

# field -> (column heading, accepted spellings). The spellings are compared
# lower-cased with spaces, dashes and underscores removed.
COLUMNS: dict[str, tuple[str, tuple[str, ...]]] = {
    "code": ("ID", ("id", "code", "templateid", "messageid", "msgid")),
    "pool": ("Pool", ("pool", "journey", "set")),
    "month": ("Month", ("month", "stage", "monthnumber")),
    "title": ("Title", ("title", "name", "subject", "headline")),
    "angle": ("Angle", ("angle", "theme")),
    "body": ("Message", ("message", "body", "text", "copy", "content", "messagetext")),
    "buckets": ("Groups", ("groups", "group", "buckets", "bucket", "replygroups")),
    "occasion": ("Occasion", ("occasion", "festival", "event")),
    "starts_on": ("From", ("from", "startson", "startdate", "start", "validfrom")),
    "ends_on": ("Until", ("until", "endson", "enddate", "end", "validuntil", "to", "expiry",
                          "expires")),
    "order_by": ("Order by", ("orderby", "cutoff", "ordercutoff", "orderingcutoff",
                              "lastorderdate")),
    "kitchens": ("Kitchens", ("kitchens", "kitchen", "servicearea", "area", "outlets")),
    "dishes": ("Dishes", ("dishes", "dish", "needsdishes", "menuitems")),
    "wa_template": ("WhatsApp template", ("whatsapptemplate", "watemplate", "template",
                                          "templatename", "gallaboxtemplate")),
    "wa_params": ("Template values", ("templatevalues", "templateparams", "waparams",
                                      "params", "values")),
    "wa_button": ("Template button", ("templatebutton", "wabutton", "button")),
    "is_active": ("Active", ("active", "isactive", "enabled", "status", "on")),
}
FIELDS = tuple(COLUMNS)
_ALIASES = {alias: name for name, (_, aliases) in COLUMNS.items() for alias in aliases}

HELP = {
    "code": "Required. A short ID that never changes, e.g. M1-01. Importing a file "
            "again updates the messages with these IDs.",
    "pool": "Which journey uses it (a follow-up campaign picks one pool). Empty: SMO.",
    "month": "The journey month it is for: 1 is the first month after the first "
             "message. Empty: any month.",
    "title": "Required. For your team; customers do not see it.",
    "angle": "The idea behind it (craving, variety...). For your team.",
    "body": f"Required. What the customer reads, up to {MESSAGE_MAX} characters. "
            "*bold* and _italic_ work. Placeholders: {first_name}, {shop_name}...",
    "buckets": "Reply groups it suits: No response, Later, No - separated by commas. "
               "Empty: all of them.",
    "occasion": "A festival or occasion name, for your team (e.g. Diwali).",
    "starts_on": "Only sent from this date (YYYY-MM-DD). With Until, an occasion "
                 "message: preferred over the month's messages between the dates.",
    "ends_on": "Only sent until this date (YYYY-MM-DD): the occasion's expiry.",
    "order_by": "The occasion's ordering cutoff (YYYY-MM-DD): not sent after this day, "
                "and the skip is logged with the reason.",
    "kitchens": "The occasion's service area: kitchen codes (as on the Kitchens page), "
                "separated by commas. Only customers of these kitchens get it. Empty: all.",
    "dishes": "Only sent while at least one of these dishes is on the menu. "
              "Separate with commas.",
    "wa_template": "The approved WhatsApp template in Gallabox with this same text, "
                   "for customers who have not written in the last 24 hours.",
    "wa_params": "Values for the template's {{1}}, {{2}}... in order, separated by |. "
                 "Usually empty.",
    "wa_button": "order_link if the template has a button to their ordering page, "
                 "else none.",
    "is_active": "Yes or No. Only active messages are sent.",
}


# --- one message's values ---------------------------------------------------------------
def _text(raw) -> str:
    if raw is None:
        return ""
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    return str(raw).replace("\r\n", "\n").replace("\r", "\n").strip()


def _one_line(raw) -> str:
    return " ".join(_text(raw).split())


def clean_code(raw) -> str:
    code = _one_line(raw).upper()
    if not code:
        raise NudgeConfigError("ID is required")
    if not CODE.match(code):
        raise NudgeConfigError(f"ID {code!r}: use up to 40 letters, digits, - _ . /")
    return code


def _month(raw) -> int | None:
    value = _one_line(raw).lower().removeprefix("month").strip()
    if value in ("", "any", "all", "-"):
        return None
    try:
        number = int(float(value))
    except ValueError:
        raise NudgeConfigError(f"Month {raw!r} is not a number") from None
    if not 1 <= number <= 24:
        raise NudgeConfigError("Month is 1 to 24, or empty for any month")
    return number


def _list(raw, *, sep: str = ",") -> list[str]:
    if raw is None:
        return []
    if isinstance(raw, (list, tuple)):
        parts = [_one_line(v) for v in raw]
    else:
        text = _text(raw)
        parts = [_one_line(v) for v in re.split(f"[{re.escape(sep)}\n;]" if sep == "," else
                                                re.escape(sep), text)]
    return [p for p in parts if p]


def _buckets(raw) -> list[str]:
    names = {}
    for bucket, label in BUCKET_LABELS.items():
        for spelling in (str(bucket), label):
            names[re.sub(r"[\s_\-]", "", spelling.lower())] = str(bucket)
    chosen = []
    for part in _list(raw):
        key = re.sub(r"[\s_\-]", "", part.lower())
        if key in ("all", "any", "everyone"):
            return []
        if key not in names:
            raise NudgeConfigError(f"Groups: {part!r} is not one of "
                                   f"{', '.join(BUCKET_LABELS.values())}")
        if names[key] not in chosen:
            chosen.append(names[key])
    # In the fixed order, and none at all when every group is ticked.
    ordered = [str(b) for b in JourneyBucket if str(b) in chosen]
    return [] if len(ordered) == len(JourneyBucket) else ordered


def _date(raw, name: str) -> date | None:
    if raw is None or raw == "":
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = _one_line(raw)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d %b %Y", "%d %B %Y", "%b %d %Y", "%B %d %Y"):
        try:
            return datetime.strptime(text.replace(",", ""), fmt).date()
        except ValueError:
            continue
    raise NudgeConfigError(f"{name} {text!r} is not a date; write it as YYYY-MM-DD")


def _bool(raw, *, default: bool = True) -> bool:
    if raw is None or raw == "":
        return default
    if isinstance(raw, bool):
        return raw
    text = _one_line(raw).lower()
    if text in ("yes", "y", "true", "1", "on", "active", "enabled"):
        return True
    if text in ("no", "n", "false", "0", "off", "inactive", "disabled", "paused"):
        return False
    raise NudgeConfigError(f"Active {raw!r}: write Yes or No")


def _button(raw) -> str:
    text = re.sub(r"[\s\-]", "_", _one_line(raw).lower()) or "none"
    if text in ("ordering_page", "order", "order_page", "link"):
        text = "order_link"
    if text not in BUTTONS:
        raise NudgeConfigError("Template button is order_link or none")
    return text


def clean(data: dict, *, existing: NudgeMessage | None = None) -> dict:
    """Validated values for the fields given (all of them for a new message).

    Raises NudgeConfigError naming the first problem.
    """
    out: dict = {}
    creating = existing is None
    if creating or "code" in data:
        out["code"] = clean_code(data.get("code"))
    if creating or "pool" in data:
        out["pool"] = _one_line(data.get("pool"))[:60] or DEFAULT_POOL
    if creating or "month" in data:
        out["month"] = _month(data.get("month"))
    if creating or "title" in data:
        title = _one_line(data.get("title"))
        if not title:
            raise NudgeConfigError("Title is required")
        if len(title) > 160:
            raise NudgeConfigError("Title is at most 160 characters")
        out["title"] = title
    if creating or "angle" in data:
        out["angle"] = _one_line(data.get("angle"))[:80] or None
    if creating or "body" in data:
        body = _text(data.get("body"))
        if not body:
            raise NudgeConfigError("Message is empty")
        if len(body) > MESSAGE_MAX:
            raise NudgeConfigError(f"Message is {len(body)} characters; at most {MESSAGE_MAX}")
        check_library_text(body, where="Message")
        out["body"] = body
    if creating or "buckets" in data:
        out["buckets"] = _buckets(data.get("buckets"))
    if creating or "occasion" in data:
        out["occasion"] = _one_line(data.get("occasion"))[:80] or None
    if creating or "starts_on" in data:
        out["starts_on"] = _date(data.get("starts_on"), "From")
    if creating or "ends_on" in data:
        out["ends_on"] = _date(data.get("ends_on"), "Until")
    if creating or "order_by" in data:
        out["order_by"] = _date(data.get("order_by"), "Order by")
    if creating or "kitchens" in data:
        kitchens = [k.upper()[:40] for k in _list(data.get("kitchens"))]
        if len(kitchens) > 50:
            raise NudgeConfigError("At most 50 kitchens")
        out["kitchens"] = kitchens
    if creating or "dishes" in data:
        dishes = _list(data.get("dishes"))
        if len(dishes) > 50:
            raise NudgeConfigError("At most 50 dishes")
        out["dishes"] = [d[:120] for d in dishes]
    if creating or "wa_template" in data:
        name = _one_line(data.get("wa_template")).lower() or None
        if name and (len(name) > 120 or not TEMPLATE_NAME.match(name)):
            raise NudgeConfigError("WhatsApp template names are lower case letters, digits and _")
        out["wa_template"] = name
    if creating or "wa_params" in data:
        params = _list(data.get("wa_params"), sep="|")
        if len(params) > 10:
            raise NudgeConfigError("At most 10 template values")
        for index, param in enumerate(params, start=1):
            check_placeholders(param, kind="follow_up", has_coupon=True,
                               where=f"Template value {index}")
        out["wa_params"] = params
    if creating or "wa_button" in data:
        out["wa_button"] = _button(data.get("wa_button"))
    if creating or "is_active" in data:
        out["is_active"] = _bool(data.get("is_active"))

    starts = out.get("starts_on", existing.starts_on if existing else None)
    ends = out.get("ends_on", existing.ends_on if existing else None)
    if starts and ends and ends < starts:
        raise NudgeConfigError("Until is before From")
    cutoff = out.get("order_by", existing.order_by if existing else None)
    if cutoff and starts and cutoff < starts:
        raise NudgeConfigError("Order by is before From")
    return out


# --- reading and saving -------------------------------------------------------------------
async def list_messages(session: AsyncSession, *, pool: str | None = None) -> list[NudgeMessage]:
    query = select(NudgeMessage)
    if pool:
        query = query.where(NudgeMessage.pool == pool)
    rows = list((await session.execute(query)).scalars())
    return sorted(rows, key=sort_key)


def sort_key(message: NudgeMessage) -> tuple:
    """Pool, then month (any-month last), then the ID with its numbers in order."""
    parts = [int(p) if p.isdigit() else p for p in re.split(r"(\d+)", message.code)]
    return (message.pool.lower(), message.month or 99, parts)


async def get_message(session: AsyncSession, message_id) -> NudgeMessage | None:
    import uuid

    try:
        return await session.get(NudgeMessage, uuid.UUID(str(message_id)))
    except (TypeError, ValueError):
        return None


async def by_code(session: AsyncSession, code: str) -> NudgeMessage | None:
    return (await session.execute(
        select(NudgeMessage).where(NudgeMessage.code == code))).scalar_one_or_none()


async def _code_free(session: AsyncSession, code: str, mine: NudgeMessage | None) -> None:
    other = await by_code(session, code)
    if other is not None and other is not mine:
        raise NudgeConfigError(f"There is already a message with the ID {code}")


async def create_message(session: AsyncSession, data: dict, *, actor: str) -> NudgeMessage:
    values = clean(_known(data))
    await _code_free(session, values["code"], None)
    message = NudgeMessage(**values, created_by=actor, updated_by=actor)
    session.add(message)
    await session.flush()
    log.info("nudge_message_created", code=message.code, actor=actor)
    return message


async def update_message(session: AsyncSession, message: NudgeMessage, data: dict, *,
                         actor: str) -> NudgeMessage:
    values = clean(_known(data), existing=message)
    if "code" in values:
        await _code_free(session, values["code"], message)
    for name, value in values.items():
        setattr(message, name, value)
    message.updated_by = actor
    await session.flush()
    log.info("nudge_message_updated", code=message.code, actor=actor, changed=sorted(values))
    return message


async def duplicate_message(session: AsyncSession, message: NudgeMessage, *,
                            actor: str) -> NudgeMessage:
    """A switched-off copy with the next free ID ("M1-01-COPY", "-COPY2"...)."""
    base = f"{message.code}-COPY"[:34]
    code, n = base, 1
    while await by_code(session, code) is not None:
        n += 1
        code = f"{base}{n}"
    values = {name: getattr(message, name) for name in FIELDS if name != "code"}
    copy = NudgeMessage(**{k: list(v) if isinstance(v, list) else v for k, v in values.items()},
                        code=code, created_by=actor, updated_by=actor)
    copy.is_active = False
    copy.title = f"Copy of {message.title}"[:160]
    session.add(copy)
    await session.flush()
    return copy


async def delete_message(session: AsyncSession, message: NudgeMessage, *, actor: str) -> None:
    """Deleted. Sends of it stay in the log under its ID."""
    log.info("nudge_message_deleted", code=message.code, actor=actor)
    await session.delete(message)
    await session.flush()


def _known(data: dict) -> dict:
    unknown = set(data) - set(FIELDS)
    if unknown:
        raise NudgeConfigError(f"{sorted(unknown)} are not message fields; use {list(FIELDS)}")
    return data


def view(message: NudgeMessage) -> dict:
    return {
        "id": str(message.id),
        "code": message.code,
        "pool": message.pool,
        "month": message.month,
        "title": message.title,
        "angle": message.angle,
        "body": message.body,
        "buckets": list(message.buckets or []),
        "occasion": message.occasion,
        "starts_on": message.starts_on.isoformat() if message.starts_on else None,
        "ends_on": message.ends_on.isoformat() if message.ends_on else None,
        "order_by": message.order_by.isoformat() if message.order_by else None,
        "kitchens": list(message.kitchens or []),
        "dishes": list(message.dishes or []),
        "wa_template": message.wa_template,
        "wa_params": list(message.wa_params or []),
        "wa_button": message.wa_button,
        "is_active": message.is_active,
        "updated_by": message.updated_by,
        "updated_at": message.updated_at.isoformat() if message.updated_at else None,
    }


# --- reading a file --------------------------------------------------------------------------
@dataclass
class FileRows:
    rows: list[tuple[int, dict]]          # (row number in the file, field -> raw value)
    columns: list[str]                     # the fields the file has
    ignored: list[str] = field(default_factory=list)   # headings not understood


def _field_for(heading) -> str | None:
    key = re.sub(r"[\s_\-.()]", "", str(heading or "").lower())
    return _ALIASES.get(key)


def _map_headings(headings: list) -> tuple[dict[int, str], list[str]]:
    mapping, ignored = {}, []
    for index, heading in enumerate(headings):
        name = _field_for(heading)
        if name and name not in mapping.values():
            mapping[index] = name
        elif str(heading or "").strip():
            ignored.append(str(heading).strip())
    if "code" not in mapping.values():
        raise NudgeConfigError("The file needs an ID column (the message's ID, e.g. M1-01)")
    return mapping, ignored


def read_file(filename: str, raw: bytes) -> FileRows:
    """The rows of an uploaded .xlsx, .csv or .json file."""
    if len(raw) > MAX_FILE_BYTES:
        raise NudgeConfigError(f"The file is over {MAX_FILE_BYTES // 1024 // 1024} MB")
    if not raw.strip():
        raise NudgeConfigError("The file is empty")
    name = (filename or "").lower()
    if name.endswith(".json") or raw.lstrip()[:1] in (b"[", b"{"):
        parsed = _read_json(raw)
    elif name.endswith((".xlsx", ".xlsm")) or raw[:2] == b"PK":
        parsed = _read_xlsx(raw)
    elif name.endswith(".xls"):
        raise NudgeConfigError("Old .xls files cannot be read; save it as .xlsx (Excel "
                               "Workbook) or CSV and try again")
    else:
        parsed = _read_csv(raw)
    if len(parsed.rows) > MAX_ROWS:
        raise NudgeConfigError(f"At most {MAX_ROWS} messages in one file")
    if not parsed.rows:
        raise NudgeConfigError("No messages found in the file")
    return parsed


def _read_json(raw: bytes) -> FileRows:
    try:
        data = json.loads(raw.decode("utf-8-sig"))
    except (UnicodeDecodeError, ValueError) as exc:
        raise NudgeConfigError(f"That is not valid JSON ({exc})") from None
    if isinstance(data, dict):
        data = data.get("messages", data.get("templates"))
    if not isinstance(data, list):
        raise NudgeConfigError('JSON must be a list of messages, or {"messages": [...]}')
    rows, columns, ignored = [], [], []
    for number, item in enumerate(data, start=1):
        if not isinstance(item, dict):
            raise NudgeConfigError(f"Message {number} in the JSON is not an object")
        values = {}
        for key, value in item.items():
            name = key if key in COLUMNS else _field_for(key)
            if name is None:
                if key not in ignored:
                    ignored.append(key)
                continue
            values[name] = value
            if name not in columns:
                columns.append(name)
        rows.append((number, values))
    if "code" not in columns:
        raise NudgeConfigError('Each message needs an "id" (e.g. "M1-01")')
    return FileRows(rows=rows, columns=columns, ignored=ignored)


def _read_csv(raw: bytes) -> FileRows:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    table = [row for row in csv.reader(io.StringIO(text), dialect)]
    return _table(table)


def _read_xlsx(raw: bytes) -> FileRows:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any unreadable workbook
        raise NudgeConfigError(f"That Excel file cannot be opened ({exc})") from None
    try:
        sheets = [s for s in book.worksheets if s.title.lower() == "messages"] or book.worksheets
        for sheet in sheets:
            table = [list(row) for row in sheet.iter_rows(values_only=True)]
            if any(_field_for(c) == "code" for row in table[:10] for c in row):
                return _table(table)
        raise NudgeConfigError("No sheet has an ID column (the message's ID, e.g. M1-01)")
    finally:
        book.close()


def _table(table: list[list]) -> FileRows:
    """Rows under the first row that has an ID heading."""
    start = next((i for i, row in enumerate(table[:10])
                  if any(_field_for(c) == "code" for c in row)), None)
    if start is None:
        raise NudgeConfigError("The file needs an ID column (the message's ID, e.g. M1-01)")
    mapping, ignored = _map_headings(table[start])
    rows = []
    for offset, row in enumerate(table[start + 1:], start=start + 2):
        values = {name: row[index] if index < len(row) else None
                  for index, name in mapping.items()}
        if all(v is None or str(v).strip() == "" for v in values.values()):
            continue
        rows.append((offset, values))
    return FileRows(rows=rows, columns=list(mapping.values()), ignored=ignored)


# --- importing --------------------------------------------------------------------------------
UPSERT, ADD_ONLY = "upsert", "add_only"
MODES = (UPSERT, ADD_ONLY)


@dataclass
class RowResult:
    row: int
    code: str
    action: str                  # create | update | unchanged | skip | error
    title: str = ""
    changes: list[str] = field(default_factory=list)
    error: str | None = None


def _comparable(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


async def import_rows(session: AsyncSession, parsed: FileRows, *, mode: str = UPSERT,
                      turn_off_missing: bool = False, commit: bool = False,
                      actor: str) -> dict:
    """Check (commit=False) or apply an import. The same report either way."""
    if mode not in MODES:
        raise NudgeConfigError(f"mode is one of {list(MODES)}")
    existing = {m.code: m for m in await list_messages(session)}
    results: list[RowResult] = []
    seen: dict[str, int] = {}
    pools: set[str] = set()
    applied = 0

    for number, raw in parsed.rows:
        try:
            code = clean_code(raw.get("code"))
        except NudgeConfigError as exc:
            results.append(RowResult(number, _one_line(raw.get("code")), "error", error=str(exc)))
            continue
        if code in seen:
            results.append(RowResult(number, code, "error",
                                     error=f"Same ID as row {seen[code]}; only the first is used"))
            continue
        seen[code] = number
        current = existing.get(code)
        try:
            if current is None:
                values = clean(raw)
            else:
                values = clean({k: v for k, v in raw.items() if k != "code"}, existing=current)
        except NudgeConfigError as exc:
            results.append(RowResult(number, code, "error", title=_one_line(raw.get("title")),
                                     error=str(exc)))
            continue

        pools.add(values.get("pool") or (current.pool if current else DEFAULT_POOL))
        if current is None:
            result = RowResult(number, code, "create", title=values["title"])
            if commit:
                session.add(NudgeMessage(**values, created_by=actor, updated_by=actor))
                applied += 1
        elif mode == ADD_ONLY:
            result = RowResult(number, code, "skip", title=current.title,
                               error="Already in the library (adding new IDs only)")
        else:
            changes = [name for name, value in values.items()
                       if _comparable(getattr(current, name)) != _comparable(value)]
            result = RowResult(number, code, "update" if changes else "unchanged",
                               title=values.get("title", current.title), changes=changes)
            if commit and changes:
                for name in changes:
                    setattr(current, name, values[name])
                current.updated_by = actor
                applied += 1
        results.append(result)

    turned_off = []
    if turn_off_missing:
        for code, message in existing.items():
            if message.pool in pools and code not in seen and message.is_active:
                turned_off.append(code)
                if commit:
                    message.is_active = False
                    message.updated_by = actor
    if commit:
        await session.flush()
        log.info("nudge_messages_imported", actor=actor, applied=applied,
                 turned_off=len(turned_off), errors=sum(r.action == "error" for r in results))

    counts = {action: sum(r.action == action for r in results)
              for action in ("create", "update", "unchanged", "skip", "error")}
    return {
        "committed": commit,
        "mode": mode,
        "counts": counts,
        "turned_off": turned_off,
        "columns": [COLUMNS[c][0] for c in parsed.columns],
        "ignored_columns": parsed.ignored,
        "rows": [r.__dict__ for r in results],
    }


# --- writing a file ----------------------------------------------------------------------------
def _row(message: NudgeMessage | dict) -> dict:
    """A message as the file's cells, the way an import reads them back."""
    get = message.get if isinstance(message, dict) else lambda k, d=None: getattr(message, k, d)
    buckets = get("buckets") or []
    return {
        "code": get("code"),
        "pool": get("pool") or DEFAULT_POOL,
        "month": get("month"),
        "title": get("title"),
        "angle": get("angle") or "",
        "body": get("body"),
        "buckets": ", ".join(BUCKET_LABELS.get(JourneyBucket(b), b) for b in buckets),
        "occasion": get("occasion") or "",
        "starts_on": get("starts_on"),
        "ends_on": get("ends_on"),
        "order_by": get("order_by"),
        "kitchens": ", ".join(get("kitchens") or []),
        "dishes": ", ".join(get("dishes") or []),
        "wa_template": get("wa_template") or "",
        "wa_params": " | ".join(get("wa_params") or []),
        "wa_button": get("wa_button") or "none",
        "is_active": "Yes" if get("is_active", True) else "No",
    }


def to_json(messages: list) -> bytes:
    items = []
    for message in messages:
        row = _row(message)
        items.append({
            "id": row["code"], "pool": row["pool"], "month": row["month"],
            "title": row["title"], "angle": row["angle"] or None, "message": row["body"],
            "groups": [BUCKET_LABELS[JourneyBucket(b)] for b in
                       ((message.get("buckets") if isinstance(message, dict)
                         else message.buckets) or [])],
            "occasion": row["occasion"] or None,
            "from": row["starts_on"].isoformat() if row["starts_on"] else None,
            "until": row["ends_on"].isoformat() if row["ends_on"] else None,
            "order_by": row["order_by"].isoformat() if row["order_by"] else None,
            "kitchens": [k.strip() for k in row["kitchens"].split(",") if k.strip()],
            "dishes": [d.strip() for d in row["dishes"].split(",") if d.strip()],
            "whatsapp_template": row["wa_template"] or None,
            "template_values": [v.strip() for v in row["wa_params"].split("|") if v.strip()],
            "template_button": row["wa_button"],
            "active": row["is_active"] == "Yes",
        })
    return json.dumps({"exported_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                       "count": len(items), "messages": items},
                      ensure_ascii=False, indent=2).encode("utf-8")


def to_csv(messages: list) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow([COLUMNS[name][0] for name in FIELDS])
    for message in messages:
        row = _row(message)
        writer.writerow([(v.isoformat() if isinstance(v, date) else ("" if v is None else v))
                         for v in (row[name] for name in FIELDS)])
    # Excel opens a UTF-8 CSV correctly only with the byte-order mark.
    return ("﻿" + out.getvalue()).encode("utf-8")


def to_xlsx(messages: list) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    sheet = book.active
    sheet.title = "Messages"
    sheet.append([COLUMNS[name][0] for name in FIELDS])
    for message in messages:
        row = _row(message)
        sheet.append([row[name] for name in FIELDS])

    widths = {"order_by": 12, "kitchens": 16,
              "code": 10, "pool": 8, "month": 8, "title": 34, "angle": 18, "body": 70,
              "buckets": 18, "occasion": 14, "starts_on": 12, "ends_on": 12, "dishes": 24,
              "wa_template": 22, "wa_params": 18, "wa_button": 15, "is_active": 8}
    head = PatternFill("solid", fgColor="028582")
    for index, name in enumerate(FIELDS, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = widths[name]
        cell = sheet.cell(row=1, column=index)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
        cell.alignment = Alignment(vertical="center")
    body_col = FIELDS.index("body") + 1
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column == body_col)
            if FIELDS[cell.column - 1] in ("starts_on", "ends_on", "order_by") and cell.value:
                cell.number_format = "yyyy-mm-dd"
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions

    last = max(sheet.max_row, 500)
    for name, choices in (("is_active", '"Yes,No"'), ("wa_button", '"none,order_link"')):
        rule = DataValidation(type="list", formula1=choices, allow_blank=True)
        letter = get_column_letter(FIELDS.index(name) + 1)
        rule.add(f"{letter}2:{letter}{last}")
        sheet.add_data_validation(rule)
    months = DataValidation(type="whole", operator="between", formula1="1", formula2="24",
                            allow_blank=True, error="Month is 1 to 24, or empty for any month")
    letter = get_column_letter(FIELDS.index("month") + 1)
    months.add(f"{letter}2:{letter}{last}")
    sheet.add_data_validation(months)

    guide = book.create_sheet("How to fill it in")
    guide.append(["Column", "What to put in it"])
    for name in FIELDS:
        guide.append([COLUMNS[name][0], HELP[name]])
    guide.column_dimensions["A"].width = 20
    guide.column_dimensions["B"].width = 110
    for cell in guide[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
    for row in guide.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


FORMATS = {
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", to_xlsx),
    "csv": ("text/csv; charset=utf-8", to_csv),
    "json": ("application/json", to_json),
}


def export(messages: list, fmt: str) -> tuple[bytes, str]:
    """(file bytes, media type) for xlsx, csv or json."""
    if fmt not in FORMATS:
        raise NudgeConfigError(f"format is one of {list(FORMATS)}")
    media_type, writer = FORMATS[fmt]
    return writer(messages), media_type


SAMPLE = [
    {"code": "M1-01", "pool": "SMO", "month": 1, "title": "Rice is sorted. What goes on the side?",
     "angle": "Meal completion",
     "body": "Rice is sorted for today's meal? 🍚\n\nNow comes the interesting part - what goes "
             "with it?\n\nExplore the curries and vegetarian sides available from Shero Home "
             "Food to find something that completes your plate.\n\nReady to explore? Send "
             "*ORDER FOOD* on WhatsApp. 🍚",
     "wa_template": "smo_m1_01", "is_active": True},
    {"code": "M2-01", "pool": "SMO", "month": 2, "title": "Kerala, Andhra or Chettinad?",
     "angle": "Regional discovery", "buckets": ["no_response", "later"],
     "body": "Hi {first_name}! Three regional food traditions. So many flavours to explore. "
             "🍚\n\nWhich would you choose?\n\nSend *ORDER FOOD* on WhatsApp to explore.",
     "wa_template": "smo_m2_01", "wa_params": ["{first_name}"], "is_active": True},
    {"code": "FEST-DIWALI", "pool": "SMO", "title": "Diwali sides for the family",
     "angle": "Festival relevance", "occasion": "Diwali",
     "starts_on": date(2026, 11, 1), "ends_on": date(2026, 11, 9),
     "body": "Celebrating Diwali with family? 🪔\n\nExplore the curries and sides available "
             "from Shero Home Food for your festive table.\n\nSend *ORDER FOOD* to begin.",
     "is_active": False},
]


def sample(fmt: str) -> tuple[bytes, str]:
    """A small file in the import layout, to start from."""
    return export(SAMPLE, fmt)
